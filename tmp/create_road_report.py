from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re, html

ROOT=Path(r'E:\FoAI_Capstone')
SRC=ROOT/'Road_Accident_Journal_Paper.tex'
OUT=ROOT/'Road_Accident_Detection_Project_Report.docx'
TMP=ROOT/'tmp'
LOGO=TMP/'amrita_logo.png'

# Reuse the institutional mark embedded in the supplied reference report.
if not LOGO.exists():
    from pypdf import PdfReader
    for page in PdfReader(r'C:\Users\gnssh\OneDrive\Desktop\reference report.pdf').pages[:3]:
        imgs=page.images
        if imgs:
            imgs[0].image.save(LOGO); break

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.85); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.95); sec.right_margin=Inches(.95)
sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); normal.font.size=Pt(12)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing=1.25; normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.first_line_indent=Inches(.25)
for nm,size in [('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
    st=styles[nm]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)
styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 2'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
styles['Heading 3'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT

# Centered page number, as in the reference.
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)

def fmt(p, align=None, first=False, before=0, after=0, line=1.0):
    if align is not None: p.alignment=align
    p.paragraph_format.first_line_indent=Inches(.25) if first else Inches(0)
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=line
    return p
def para(text='',bold=False,italic=False,size=12,align=None,first=False,before=0,after=0,line=1.0):
    p=fmt(doc.add_paragraph(),align,first,before,after,line); run=p.add_run(text); run.bold=bold; run.italic=italic; run.font.name='Times New Roman'; run.font.size=Pt(size); return p
def centered(text,size=14,bold=True,italic=False,before=0,after=12):
    return para(text,bold,italic,size,WD_ALIGN_PARAGRAPH.CENTER,False,before,after,1.0)
def pagebreak(): doc.add_page_break()
def logo_line():
    p=fmt(doc.add_paragraph(),WD_ALIGN_PARAGRAPH.CENTER,False,0,16,1.0)
    if LOGO.exists(): p.add_run().add_picture(str(LOGO),width=Inches(4.55))

title='AI-BASED ROAD ACCIDENT DETECTION, SEVERITY ASSESSMENT AND EMERGENCY ALERT SYSTEM FOR NEARBY HOSPITALS'
authors=['PAVAN ADITYA','G SRI HARSHA','M PALLAVI']

# PAGE 1 — reference cover structure.
for _ in range(2): para('',after=8)
for line in ['AI-BASED ROAD ACCIDENT DETECTION,','SEVERITY ASSESSMENT AND EMERGENCY','ALERT SYSTEM FOR NEARBY HOSPITALS']:
    centered(line,16,True,False,0,8)
centered('A PROJECT REPORT',14,True,False,12,10)
centered('Submitted to',12,True,True,0,10)
centered('Amrita Vishwa Vidyapeetham',14,True,False,0,12)
centered('in partial fulfillment for the award of the degree of',12,True,True,0,10)
centered('BACHELOR OF TECHNOLOGY IN COMPUTER SCIENCE AND',14,True,False,0,4)
centered('ENGINEERING',14,True,False,0,13)
centered('by',12,True,True,0,10)
for name in authors:
    centered(name,13,True,False,0,2)
    centered('[Register number not provided in project files]',10,False,False,0,8)
centered('Supervisor',12,True,True,3,7)
centered('[Supervisor name and designation not provided]',12,True,False,0,8)
logo_line()
centered('AMRITA VISHWA VIDYAPEETHAM',14,True,False,0,7)
centered('AMRITA SCHOOL OF COMPUTING',14,True,False,0,7)
centered('CHENNAI – 601103',13,True,False,0,9)
centered('OCTOBER 2026',13,True,False,0,0)

# PAGE 2 — bonafide certificate.
pagebreak(); logo_line(); centered('BONAFIDE CERTIFICATE',14,True,False,0,24)
para('This is to certify that this project report entitled “'+title.title()+'” is the bonafide work of “Pavan Aditya”, “G Sri Harsha”, and “M Pallavi”, submitted in partial fulfillment of the requirements for the award of Bachelor of Technology in Computer Science and Engineering at Amrita Vishwa Vidyapeetham.',first=False,after=20,line=1.5)
para('The project supervisor’s name and designation, candidate register numbers, and institutional signatories were not present in the supplied project files and must be completed by the department before formal submission.',first=False,after=30,line=1.5)
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for c,txt in zip(t.rows[0].cells,['SIGNATURE\n\n[Project supervisor]\nSupervisor\nDepartment of CSE\nAmrita School of Computing\nChennai','SIGNATURE\n\n[Department chairperson]\nChairperson\nDepartment of CSE\nAmrita School of Computing\nChennai']):
    c.width=Inches(3.0); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; c.text=txt
    for p in c.paragraphs:
        p.paragraph_format.space_after=Pt(10)
        for rr in p.runs: rr.font.name='Times New Roman'; rr.font.size=Pt(11); rr.bold=True
para('',after=45)
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,txt in zip(t.rows[0].cells,['INTERNAL EXAMINER','EXTERNAL EXAMINER']):
    c.text='\n\n'+txt
    for rr in c.paragraphs[0].runs: rr.font.name='Times New Roman'; rr.font.size=Pt(11); rr.bold=True

# PAGE 3 — declaration.
pagebreak(); logo_line(); para('',after=18); centered('DECLARATION BY THE CANDIDATES',14,True,False,0,25)
para('We declare that the report entitled “'+title.title()+'” submitted by us for the degree of Bachelor of Technology is the record of the project work carried out by us. The project files identify the student authors and the project implementation, but do not identify the project supervisor or provide candidate register numbers. These details are therefore left for verification rather than assumed.',first=False,after=18,line=1.5)
para('This work has not formed the basis for the award of any degree, diploma, associateship, fellowship, or title in this or any other University or similar institution of higher learning.',first=False,after=25,line=1.5)
for name in authors:
    p=fmt(doc.add_paragraph(),WD_ALIGN_PARAGRAPH.RIGHT,False,2,8,1.0)
    rr=p.add_run('SIGNATURE\n'+name+'\n(Register No. not supplied)'); rr.bold=True; rr.font.name='Times New Roman'; rr.font.size=Pt(11)

# PAGE 4 — abstract.
pagebreak(); para('',after=45); centered('ABSTRACT',14,True,False,0,20)
abstract=('This project presents a prototype for road accident monitoring that combines computer-vision detection, visual severity labels, video-level processing, and an emergency-alert workflow. The supplied repository contains three YOLO26n experiments: a ten-class accident-interaction detector, a three-class severity detector, and a five-class night-traffic road-user detector. Saved final validation records report precision, recall, mAP@0.5, and mAP@0.5:0.95 of 94.701%, 91.308%, 97.321%, and 84.124% for accident detection; 76.174%, 77.897%, 82.157%, and 60.530% for severity; and 83.960%, 77.684%, 85.307%, and 57.371% for night traffic. A separate saved severity test summary reports 79.74% precision, 71.81% recall, 80.45% mAP@0.5, and 55.47% mAP@0.5:0.95. The report examines the datasets and saved training curves, the video-processing logic, the FastAPI services, location lookup, and alert-state handling. Source inspection shows that spatial and scene-context checks do not currently gate event confirmation, one hospital fallback is independent of the supplied coordinates, and alert records are held in process memory without inspected authentication. The reported scores are saved project artifacts and have not been independently reproduced. They do not establish event-level reliability, injury assessment, emergency-service integration, or operational readiness. The system is therefore documented as a human-supervised academic prototype, with data provenance, camera-disjoint evaluation, secure alert handling, and partner-led validation identified as future work.')
para(abstract,first=False,after=22,line=1.5)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(9)
r=p.add_run('Keywords: '); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12)
r=p.add_run('road accident detection, computer vision, YOLO26, visual severity, video analytics, temporal verification, emergency alert workflow'); r.font.name='Times New Roman'; r.font.size=Pt(12)

# PAGE 5 — acknowledgement in the reference placement and hierarchy.
pagebreak(); centered('ACKNOWLEDGEMENT',14,True,False,0,18)
ack=[
'The project team acknowledges Amrita Vishwa Vidyapeetham and the Amrita School of Computing for the academic setting in which this capstone work was undertaken.',
'The supplied project files do not record the supervisor, project coordinator, reviewers, or other individual contributors by name. No individual acknowledgement has therefore been attributed without supporting information. The authors may complete this page after confirming the people and roles who directly supported the work.',
'The team also acknowledges the maintainers of the open-source software, research datasets, and technical documentation used in the project, as cited in the report references. The saved training outputs, source code, and implementation materials supplied with the project formed the basis for documenting the prototype and its limitations.'
]
for s in ack: para(s,first=True,after=9,line=1.45)
for name in authors: para(name,bold=True,first=False,after=1,line=1.0)

# The reference's contents and lists begin after the required five-page front matter.
pagebreak(); centered('TABLE OF CONTENTS',14,True,False,0,18)
toc=[
('Abstract','4'),('Acknowledgement','5'),('List of Tables','8'),('List of Figures','8'),('List of Abbreviations','9'),
('CHAPTER 1  INTRODUCTION','10'),('1.1 Background and Motivation','10'),('1.2 Problem Definition and Evaluation Objectives','11'),('1.3 Project Aim and Objectives','12'),('1.4 Scope and Boundaries','13'),('1.5 Organization of the Report','14'),
('CHAPTER 2  LITERATURE REVIEW','15'),('2.1 Accident Detection and Temporal Evidence','15'),('2.2 Object Detection for Road Scenes','16'),('2.3 Severity Assessment and Response Coordination','17'),('2.4 Research Gap and Project Positioning','18'),('2.5 Summary','19'),
('CHAPTER 3  METHODOLOGY AND SYSTEM ARCHITECTURE','20'),('3.1 Requirements and Design Goals','20'),('3.2 System Overview and Data Flow','21'),('3.3 Detection and Severity Models','22'),('3.4 Temporal Verification and Context Checks','23'),('3.5 Application and Alert Workflow','24'),('3.6 Summary','25'),
('CHAPTER 4  IMPLEMENTATION AND EXPERIMENTAL METHOD','26'),('4.1 Software and System Components','26'),('4.2 Dataset Inventory and Labels','27'),('4.3 Training Configuration','29'),('4.4 Video Pipeline and APIs','30'),('4.5 Summary','31'),
('CHAPTER 5  RESULTS, DISCUSSION AND LIMITATIONS','32'),('5.1 Evaluation Metrics','32'),('5.2 Detection Results','33'),('5.3 Severity and Night Traffic Results','35'),('5.4 Integration Audit','37'),('5.5 Threats to Validity','39'),
('CHAPTER 6  CONCLUSION AND FUTURE WORK','40'),('References','42')]
for label,num in toc:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18 if label[:1].isdigit() else 0); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.1))
    rr=p.add_run(label+'\t'+num); rr.font.name='Times New Roman'; rr.font.size=Pt(10.5); rr.bold=label.startswith('CHAPTER')

pagebreak(); centered('LIST OF FIGURES',14,True,False,0,16)
figlist=[('Figure 5.1','Accident detector training and validation history'),('Figure 5.2','Severity detector training and validation history'),('Figure 5.3','Night-traffic detector training and validation history'),('Figure 5.4','Accident detector validation confusion matrix'),('Figure 5.5','Severity detector validation confusion matrix'),('Figure 5.6','Supplied project comparison graphic')]
for a,b in figlist: para(a+'  '+b,first=False,after=4,line=1.0)
centered('LIST OF TABLES',14,True,False,14,12)
for s in ['Table 3.1  Prototype processing stages and evidence basis','Table 4.1  Dataset inventory reported in supplied project artifacts','Table 4.2  Training settings identified in saved arguments','Table 4.3  Prototype functions and evidence boundaries','Table 5.1  Final-epoch validation metrics saved with the project runs','Table 5.2  Saved severity test results by class']:
    para(s,first=False,after=4,line=1.0)

pagebreak(); centered('LIST OF ABBREVIATIONS',14,True,False,0,18)
abbr=[('API','Application Programming Interface'),('AP','Average Precision'),('CCTV','Closed-Circuit Television'),('CORS','Cross-Origin Resource Sharing'),('GPU','Graphics Processing Unit'),('IoU','Intersection over Union'),('mAP','Mean Average Precision'),('NMS','Non-Maximum Suppression'),('OSM','OpenStreetMap'),('PR','Precision–Recall'),('REST','Representational State Transfer'),('TP','True Positive'),('FP','False Positive'),('FN','False Negative'),('YOLO','You Only Look Once')]
for a,b in abbr:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.35)); rr=p.add_run(a+'\t'+b); rr.font.name='Times New Roman'; rr.font.size=Pt(12)

# Convert the existing evidence-based manuscript into the six-chapter structure.
raw=SRC.read_text(encoding='utf-8')
bibs={k:v for k,v in re.findall(r'\\bibitem\{([^}]+)\}\s*(.*?)(?=\\bibitem\{|\\end\{thebibliography\})',raw,re.S)}
keys=list(bibs); refnum={k:i+1 for i,k in enumerate(keys)}
body=raw.split('\\begin{document}',1)[1].split('\\end{document}',1)[0]
body=re.sub(r'\\maketitle.*?\\clearpage','',body,flags=re.S)
body=re.sub(r'\\begin\{abstract\}.*?\\end\{abstract\}','',body,flags=re.S)
body=re.sub(r'\\tableofcontents|\\listoffigures|\\listoftables|\\pagenumbering\{arabic\}|\\clearpage|\\appendix','',body)
body=re.sub(r'\\begin\{thebibliography\}.*?\\end\{thebibliography\}','',body,flags=re.S)
# Build topic-to-chapter map from source manuscript section headings.
chapter={
'Introduction':1,'Problem Definition and Evaluation Objectives':1,'Project Scope, Requirements, and Design Goals':1,'Report Scope and Evidence Conventions':1,
'Related Work':2,'System Architecture':3,'Architecture and Data Flow in More Detail':3,
'Dataset and Experimental Method':4,'Integrated Video Pipeline: Code-Level Analysis':4,'Application, API, and Emergency-Workflow Audit':4,
'Results':5,'Data Inventory and Label Interpretation':5,'Detailed Model Results and Interpretation':5,'Limitations and Threats to Validity':5,
'Conclusion':6,'Recommended Project Work Plan':6,'Reproducibility and Artifact Checklist':6,
'Supplementary Saved Run Figures':0,'Accident Detection Training and Validation Previews':0,'Accident Severity Training and Validation Previews':0,'Night Traffic Training and Validation Previews':0}
sections=re.split(r'\\(section|subsection|subsubsection)\{([^}]+)\}',body)
current_ch=0; heading_counts={}; fig_count=0; table_count=0; inserted=set()
def clean_latex(s):
    s=re.sub(r'\\cite\{([^}]+)\}',lambda m:', '.join('['+str(refnum.get(k,k)) +']' for k in m.group(1).split(',')),s)
    s=re.sub(r'\\(?:textbf|emph|textit|texttt)\{([^{}]*)\}',r'\1',s)
    s=re.sub(r'\\url\{([^}]+)\}',r'\1',s)
    s=s.replace('\\%', '%').replace('\\&','&').replace('\\_','_').replace('\\#','#')
    s=s.replace('---','—').replace('--','–').replace('\\times','×').replace('\\rightarrow','→')
    s=re.sub(r'\\[a-zA-Z]+\*?(?:\[[^]]*\])?', ' ', s)
    s=s.replace('{','').replace('}','')
    s=re.sub(r'\$([^$]+)\$',r'\1',s)
    s=re.sub(r'\s+',' ',s)
    return html.unescape(s).strip()
def add_chapter(n,title):
    p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.keep_with_next=True; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(22)
    rr=p.add_run(f'CHAPTER {n}\n{title.upper()}'); rr.bold=True; rr.font.name='Times New Roman'; rr.font.size=Pt(14)
def add_table(data,caption):
    table_counts[current_ch]=table_counts.get(current_ch,0)+1
    para(f'Table {current_ch}.{table_counts[current_ch]} {caption}',bold=True,size=10.5,align=WD_ALIGN_PARAGRAPH.CENTER,first=False,after=5)
    t=doc.add_table(rows=0,cols=len(data[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for ri,row in enumerate(data):
        cells=t.add_row().cells
        for c,val in zip(cells,row):
            c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(9); r.bold=ri==0
    para('',after=5)
def add_figure(path,caption,width=5.8):
    path=ROOT/path
    if not path.exists(): return
    figure_counts[current_ch]=figure_counts.get(current_ch,0)+1
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3); p.paragraph_format.keep_together=True
    p.add_run().add_picture(str(path),width=Inches(width))
    para(f'Figure {current_ch}.{figure_counts[current_ch]} {caption}',size=10.5,align=WD_ALIGN_PARAGRAPH.CENTER,first=False,after=10,line=1.0)

# Group manuscript sections into the reference's chapter order. The source paper
# interleaves implementation audits, results, and limitations, so preserve its
# content while placing each topic in its report chapter.
blocks=[]
active_chapter=0
for i in range(0,len(sections)-2,3):
    kind,title=sections[i+1],sections[i+2]
    chunk=sections[i+3] if i+3<len(sections) else ''
    if kind=='section':
        active_chapter=chapter.get(title.strip(),0)
    if title.strip() and active_chapter:
        blocks.append((kind,title.strip(),chunk,active_chapter))

chnames={1:'INTRODUCTION',2:'LITERATURE REVIEW',3:'METHODOLOGY AND SYSTEM ARCHITECTURE',4:'IMPLEMENTATION AND EXPERIMENTAL METHOD',5:'RESULTS, DISCUSSION AND LIMITATIONS',6:'CONCLUSION AND FUTURE WORK'}
chapter_sections={1:0,2:0,3:0,4:0,5:0,6:0}; subsection_counts={}; figure_counts={}; table_counts={}
# Manually authored scope/objectives are based on the supplied manuscript and implementation.
for n in range(1,7):
  current_ch=n; add_chapter(n,chnames[n]); chapter_sections[n]=0
  for kind,title,chunk,block_chapter in blocks:
    if block_chapter!=n: continue
    if kind=='section':
        chapter_sections[n]+=1; subsection_counts[(n,chapter_sections[n])]=0
        if title not in ['Introduction','Related Work','System Architecture','Dataset and Experimental Method','Results','Conclusion']:
            h=doc.add_paragraph(style='Heading 2'); h.paragraph_format.keep_with_next=True
            h.add_run(f'{n}.{chapter_sections[n]}  {title}')
    elif kind in ('subsection','subsubsection'):
        parent=chapter_sections[n]; subsection_counts[(n,parent)]=subsection_counts.get((n,parent),0)+1
        h=doc.add_paragraph(style='Heading 2' if kind=='subsection' else 'Heading 3')
        h.add_run(f'{n}.{parent}.{subsection_counts[(n,parent)]}  {title}')
    # strip floats and environments; include stable summary tables/figures at the relevant chapter locations below.
    chunk=re.sub(r'\\begin\{(?:figure|table|tabular|minipage)\}.*?\\end\{(?:figure|table|tabular|minipage)\}',' ',chunk,flags=re.S)
    chunk=re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}',lambda m:'\nEQUATION: '+clean_latex(m.group(1))+'\n',chunk,flags=re.S)
    chunk=re.sub(r'\\label\{[^}]+\}|\\ref\{([^}]+)\}',lambda m:'[see '+(m.group(1) or '')+']' if m.group(1) else '',chunk)
    chunk=re.sub(r'\\begin\{[^}]+\}|\\end\{[^}]+\}', ' ',chunk)
    chunk=chunk.replace('\\toprule',' ').replace('\\midrule',' ').replace('\\bottomrule',' ')
    chunk=re.sub(r'\\caption\{([^}]+)\}',r' ',chunk)
    chunk=re.sub(r'\\includegraphics(?:\[[^]]*\])?\{[^}]+\}',' ',chunk)
    paras=re.split(r'\n\s*\n',chunk)
    for rawp in paras:
        txt=clean_latex(rawp)
        if len(txt)<45 or txt.startswith('EQUATION:') and len(txt)<15: continue
        if re.fullmatch(r'[\s\d&|\\.\-]+',txt): continue
        para(txt,first=True,after=6,line=1.25)
    # Project evidence figures and structured tables, placed with matching chapters.
    if title=='Processing stages' and 'architecture' not in inserted:
        inserted.add('architecture')
        add_table([['Stage','Input / output','Evidence status'],['Video preparation','Video frames to image arrays','Described in source and demo records'],['Object detection','Frames to boxes and class scores','Three saved YOLO26n runs'],['Event verification','Frame detections to candidate event','Code reviewed; checks currently non-gating'],['Location and alert','Event summary to facility and alert record','Prototype services; no dispatch validation']], 'Prototype processing stages and evidence basis.')
    if title=='Dataset and Experimental Method' and 'data' not in inserted:
        inserted.add('data')
        add_table([['Experiment','Images (train/val/test)','Labels'],['Accident detection','5,849 / 712 / 1,069 (7,630 total)','10 classes'],['Severity','3,559 / 447 / 446 (4,452 total)','Moderate, no accident, severe'],['Night traffic','Counts not identified in supplied report text','5 road-user classes']], 'Dataset inventory reported in supplied project artifacts.')
        add_table([['Parameter','Saved run setting'],['Model family','YOLO26n'],['Epochs','100'],['Image size','640 pixels'],['Batch size','8'],['Optimizer','AdamW'],['Seed','0; deterministic mode recorded'],['Patience','20'],['Augmentations','Mosaic, horizontal flip, scale, HSV; Mosaic disabled for final 10 epochs']], 'Training settings identified in saved arguments.')
    if title=='Results' and 'metrics' not in inserted:
        inserted.add('metrics')
        add_table([['Task','Precision %','Recall %','mAP@0.5 %','mAP@0.5:0.95 %'],['Accident detection','94.701','91.308','97.321','84.124'],['Severity','76.174','77.897','82.157','60.530'],['Night traffic','83.960','77.684','85.307','57.371']], 'Final-epoch validation metrics saved with the project runs.')
        for model,path in [('Accident','results/detection/yolo26_baseline/results.png'),('Severity','results/detection/accident_severity-2/results.png'),('Night traffic','results/detection/night_traffic-2/results.png')]:
            add_figure(Path(path),model+' model training and validation history.',5.7)
        add_table([['Severity class','Test precision %','Test recall %','Test AP@0.5 %'],['Moderate accident','76.42','74.65','81.39'],['No accident','80.45','70.30','81.43'],['Severe accident','82.36','70.50','78.54']], 'Saved severity test results by class.')
        add_figure(Path('results/detection/yolo26_baseline/confusion_matrix.png'),'Accident detector validation confusion matrix.',4.8)
        add_figure(Path('results/detection/accident_severity-2/confusion_matrix.png'),'Severity detector validation confusion matrix.',4.8)
        p=para('The supplied comparison graphic is retained as a project artifact. Its external method labels and reported scores are not treated as a matched benchmark because the underlying protocols and source records are not established in the supplied material.',first=True,after=5,line=1.25)
        add_figure(Path('results/project_comparison.png'),'Supplied comparison graphic; cross-study values are shown for context only and are not independently verified.',5.8)
    if title=='Application, API, and Emergency-Workflow Audit' and 'workflow_table' not in inserted:
        inserted.add('workflow_table')
        add_table([['Function','Evidence in supplied project','Boundary'],['Video inference','Saved models and demo/video code','No complete end-to-end benchmark'],['Temporal confirmation','Three-frame confirmation and 30-frame gap logic','Spatial/context checks do not gate confirmation'],['Facility lookup','Overpass query, distance sorting, cache, fixed fallbacks','One fallback ignores query coordinates'],['Alert workflow','Create, acknowledge, and dispatch-status endpoints','In-memory state; inspected endpoints lack authentication'],['Dashboard','FastAPI service and monitoring views','No verified live dispatch integration']], 'Prototype functions and evidence boundaries.')

# References use the numbered IEEE-like style already present in the source manuscript.
doc.add_page_break(); p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('REFERENCES'); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(14)
for k in keys:
    txt=clean_latex(bibs[k])
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.28); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.0
    rr=p.add_run(f'[{refnum[k]}] {txt}'); rr.font.name='Times New Roman'; rr.font.size=Pt(10)

# Metadata and refresh-on-open support for page fields.
doc.core_properties.title=title.title()
doc.core_properties.subject='Academic project report based on supplied road accident detection code and experiment artifacts'
doc.core_properties.author='Pavan Aditya; G Sri Harsha; M Pallavi'
settings=doc.settings._element
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
doc.save(OUT)
print(OUT)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'figures',sum(figure_counts.values()),'references',len(keys))
